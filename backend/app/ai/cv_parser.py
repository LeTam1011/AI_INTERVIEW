"""CV text extraction and deterministic profile normalization.

This module intentionally does not call an external AI service. The profile is a
first-pass heuristic structure that later stages can enrich with an LLM/RAG model.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


@dataclass
class ParsedCV:
    text: str
    profile: dict[str, Any]


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from a PDF or DOCX, rejecting unsupported/empty documents."""
    if not content:
        raise ValueError("The uploaded file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds the 5 MB upload limit")

    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload a PDF or DOCX file")

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported")
            chunks = [(page.extract_text() or "") for page in reader.pages]
            text = "\n".join(chunks)
        else:
            document = Document(BytesIO(content))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join(paragraphs + table_text)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this document. Check that it is a valid PDF or DOCX") from exc

    text = normalize_text(text)
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS]
    if len(text.strip()) < 30:
        raise ValueError("No usable text was found. The CV may be scanned; OCR is not supported yet")
    return text


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[\t\f\v]+", " ", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


SKILL_TERMS = [
    "Python", "Java", "JavaScript", "TypeScript", "Kotlin", "C++", "C#", "PHP", "Go", "Rust",
    "React", "ReactJS", "Angular", "Vue", "Node.js", "FastAPI", "Django", "Flask", "Spring Boot",
    "HTML", "CSS", "Tailwind", "SQL", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis",
    "Docker", "Kubernetes", "Git", "GitHub", "REST API", "GraphQL", "Firebase", "AWS", "Azure",
    "Machine Learning", "Deep Learning", "LLM", "RAG", "ChromaDB", "Pandas", "NumPy", "TensorFlow", "PyTorch",
    "OOP", "Data Structures", "Algorithms", "Linux", "Figma", "Jira", "Android", "Jetpack Compose",
]

SECTION_ALIASES = {
    "education": ["education", "học vấn", "đào tạo", "academic"],
    "experience": ["experience", "work experience", "kinh nghiệm", "employment", "thực tập"],
    "projects": ["projects", "dự án", "project experience", "personal projects"],
    "certifications": ["certifications", "certificates", "chứng chỉ", "chứng nhận"],
}


def _section_content(text: str, aliases: list[str]) -> str:
    lines = text.splitlines()
    start = None
    for index, line in enumerate(lines):
        normalized = line.strip().lower().strip(" :|-")
        if any(normalized == alias or normalized.startswith(alias + " ") for alias in aliases):
            start = index + 1
            break
    if start is None:
        return ""

    end_markers = [alias for values in SECTION_ALIASES.values() for alias in values if alias not in aliases]
    selected = []
    for line in lines[start:]:
        normalized = line.strip().lower().strip(" :|-")
        if any(normalized == marker or normalized.startswith(marker + " ") for marker in end_markers):
            break
        if line.strip():
            selected.append(line.strip())
    return "\n".join(selected)[:6000]


def build_candidate_profile(text: str) -> dict[str, Any]:
    """Create a transparent, rule-based candidate profile from extracted CV text."""
    clean = normalize_text(text)
    email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", clean, re.I)
    phone_match = re.search(r"(?<!\w)(?:\+?\d[\d .()\-]{7,}\d)(?!\w)", clean)
    links = list(dict.fromkeys(re.findall(r"(?:https?://|www\.)[^\s<>]+", clean, re.I)))[:10]

    first_lines = [line.strip() for line in clean.splitlines() if line.strip()]
    name = None
    for line in first_lines[:5]:
        if len(line) <= 70 and not re.search(r"@|https?://|\d{4,}|curriculum vitae|resume|\bCV\b", line, re.I):
            if len(line.split()) <= 6:
                name = line
                break

    lowered = clean.casefold()
    skills = []
    for skill in SKILL_TERMS:
        pattern = r"(?<![\w+#.])" + re.escape(skill.casefold()) + r"(?![\w+#.])"
        if re.search(pattern, lowered):
            skills.append(skill)

    sections = {
        key: _section_content(clean, aliases)
        for key, aliases in SECTION_ALIASES.items()
    }
    # Keep evidence snippets instead of claiming inferred facts that are not present.
    experience_years = None
    year_matches = [int(year) for year in re.findall(r"\b(19\d{2}|20\d{2})\b", clean)]
    if len(year_matches) >= 2:
        span = max(year_matches) - min(year_matches)
        if 0 <= span <= 50:
            experience_years = span

    return {
        "profile_version": 1,
        "extraction_method": "rule_based",
        "name": name,
        "email": email_match.group(0) if email_match else None,
        "phone": re.sub(r"\s+", " ", phone_match.group(0)).strip() if phone_match else None,
        "links": links,
        "skills": skills,
        "education_text": sections["education"],
        "experience_text": sections["experience"],
        "projects_text": sections["projects"],
        "certifications_text": sections["certifications"],
        "estimated_year_span": experience_years,
        "text_length": len(clean),
        "limitations": [
            "Profile fields are extracted with deterministic rules, not an LLM.",
            "Scanned PDFs require OCR, which is not implemented.",
            "Estimated year span is not a verified measure of professional experience.",
        ],
    }


def parse_cv(filename: str, content: bytes) -> ParsedCV:
    text = extract_text(filename, content)
    return ParsedCV(text=text, profile=build_candidate_profile(text))
